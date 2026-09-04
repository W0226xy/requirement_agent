from io import BytesIO

import pymupdf
from docx import Document

from requirement_agent.parsers.document import DocxParser, PdfParser
from requirement_agent.parsers.image import ImageParser


class FakeOcrEngine:
    def extract_text(self, content: bytes) -> list[str]:
        assert content == b"image-content"
        return ["Requirement title", "Acceptance criterion"]


async def test_pdf_parser_extracts_text() -> None:
    document = pymupdf.open()
    page = document.new_page()
    page.insert_text((72, 72), "PDF requirement")
    content = document.tobytes()
    document.close()

    result = await PdfParser().parse(content)

    assert result.parsed_text is not None
    assert "PDF requirement" in result.parsed_text


async def test_docx_parser_extracts_paragraphs_and_tables() -> None:
    document = Document()
    document.add_paragraph("DOCX requirement")
    table = document.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "Acceptance criterion"
    output = BytesIO()
    document.save(output)

    result = await DocxParser().parse(output.getvalue())

    assert result.parsed_text == "DOCX requirement\nAcceptance criterion"


async def test_image_parser_uses_replaceable_ocr_engine() -> None:
    result = await ImageParser(FakeOcrEngine()).parse(b"image-content")

    assert result.ocr_text == "Requirement title\nAcceptance criterion"
