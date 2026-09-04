import asyncio
from io import BytesIO

import pymupdf
from docx import Document

from requirement_agent.parsers.base import ParsedAttachment
from requirement_agent.shared.errors import AttachmentProcessingError

PDF_TYPE = "application/pdf"
DOCX_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


class PdfParser:
    def supports(self, file_type: str) -> bool:
        return file_type == PDF_TYPE

    async def parse(self, content: bytes) -> ParsedAttachment:
        return await asyncio.to_thread(self._parse, content)

    def _parse(self, content: bytes) -> ParsedAttachment:
        try:
            with pymupdf.open(  # type: ignore[no-untyped-call]
                stream=content,
                filetype="pdf",
            ) as document:
                if document.needs_pass:
                    raise AttachmentProcessingError("password-protected PDFs are not supported")
                text = "\n".join(page.get_text("text") for page in document).strip()
            return ParsedAttachment(parsed_text=text)
        except AttachmentProcessingError:
            raise
        except Exception as exc:
            raise AttachmentProcessingError("failed to parse PDF") from exc


class DocxParser:
    def supports(self, file_type: str) -> bool:
        return file_type == DOCX_TYPE

    async def parse(self, content: bytes) -> ParsedAttachment:
        return await asyncio.to_thread(self._parse, content)

    def _parse(self, content: bytes) -> ParsedAttachment:
        try:
            document = Document(BytesIO(content))
            paragraphs = [paragraph.text for paragraph in document.paragraphs if paragraph.text]
            table_cells = [
                cell.text
                for table in document.tables
                for row in table.rows
                for cell in row.cells
                if cell.text
            ]
            return ParsedAttachment(parsed_text="\n".join([*paragraphs, *table_cells]).strip())
        except Exception as exc:
            raise AttachmentProcessingError("failed to parse DOCX") from exc
